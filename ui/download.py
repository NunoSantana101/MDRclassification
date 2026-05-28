"""Download bundle: Word document + JSON files in a single ZIP.

Word document styling is modelled on the saimoneTauri pdf_sidecar.py
DOCX generator (sAImoneTitle / H1-H3 / Body palette).
"""

from __future__ import annotations

import io
import json
import re
import zipfile
from datetime import datetime, timezone


_TEAL = (0x1e, 0x8c, 0x8c)
_NAVY = (0x1a, 0x3a, 0x5c)
_GREY = (0x88, 0x88, 0x88)


def _add_styles(doc):
    from docx.shared import Pt, RGBColor
    from docx.enum.style import WD_STYLE_TYPE

    styles = doc.styles

    def _ensure(name, font_size=10, bold=False, italic=False,
                color=None, space_after=6, space_before=0, font_name="Calibri"):
        try:
            s = styles[name]
        except KeyError:
            s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.font.name = font_name
        s.font.size = Pt(font_size)
        s.font.bold = bold
        s.font.italic = italic
        if color:
            s.font.color.rgb = RGBColor(*color)
        s.paragraph_format.space_after = Pt(space_after)
        s.paragraph_format.space_before = Pt(space_before)
        return s

    _ensure("MDRTitle", font_size=22, bold=True, color=_NAVY, space_after=4)
    _ensure("MDRSubtitle", font_size=11, italic=True, color=_TEAL, space_after=12)
    _ensure("MDRH1", font_size=16, bold=True, color=_NAVY, space_before=12, space_after=6)
    _ensure("MDRH2", font_size=13, bold=True, color=_TEAL, space_before=10, space_after=4)
    _ensure("MDRH3", font_size=11, bold=True, color=_NAVY, space_before=8, space_after=4)
    _ensure("MDRBody", font_size=10, space_after=6)
    _ensure("MDRQuote", font_size=10, italic=True, color=(0x33, 0x33, 0x33), space_after=6)
    _ensure("MDRDisclaimer", font_size=8, italic=True, color=_GREY, space_after=0, space_before=12)


def _add_inline(paragraph, text: str) -> None:
    """Render simple **bold** and *italic* markdown inside a paragraph."""
    if not text:
        return
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def _add_kv_table(doc, rows: list[tuple[str, str]]) -> None:
    from docx.shared import Pt

    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Light Grid Accent 1"
    for i, (k, v) in enumerate(rows):
        c1 = table.cell(i, 0)
        c2 = table.cell(i, 1)
        c1.text = k
        c2.text = v if v else "—"
        for cell in (c1, c2):
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
        for run in c1.paragraphs[0].runs:
            run.bold = True
    doc.add_paragraph()


def _yn(v) -> str:
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return v or "—"


def _lst(v) -> str:
    if not v:
        return "—"
    return ", ".join(v) if isinstance(v, list) else str(v)


def _render_device_parameters(doc, p: dict) -> None:
    """Render the seven sections of structured device parameters."""
    doc.add_paragraph("Device Parameters", style="MDRH1")

    doc.add_paragraph("1. Identity and purpose", style="MDRH2")
    _add_kv_table(doc, [
        ("Description", p.get("device_description", "")),
        ("Intended purpose (medical claim)", p.get("intended_purpose", "")),
        ("Reusability", p.get("reusability", "")),
        ("Sterile state", p.get("sterile_state", "")),
        ("Measuring function", _yn(p.get("measuring_function"))),
    ])

    doc.add_paragraph("2. Invasiveness and body contact", style="MDRH2")
    _add_kv_table(doc, [
        ("Invasiveness category", p.get("invasiveness_category", "")),
        ("Non-invasive contact type", p.get("non_invasive_contact_type", "") or "—"),
        ("Anatomical contact site(s)", _lst(p.get("anatomical_contact_sites"))),
        ("Connected to an active device", _yn(p.get("connected_to_active_device"))),
    ])

    doc.add_paragraph("3. Duration of use", style="MDRH2")
    _add_kv_table(doc, [
        ("Continuous use duration", p.get("duration_of_use", "")),
    ])

    doc.add_paragraph("4. Active device function", style="MDRH2")
    active_rows: list[tuple[str, str]] = [
        ("Active device", _yn(p.get("is_active_device"))),
    ]
    if p.get("is_active_device") == "Yes":
        active_rows.extend([
            ("Active function(s)", _lst(p.get("active_functions"))),
            ("Potentially hazardous energy administration",
             _yn(p.get("hazardous_energy_administration"))),
            ("Monitors vital parameters (immediate danger)",
             _yn(p.get("monitors_vital_immediate_danger"))),
            ("Integrated closed-loop diagnostic (Rule 22)",
             _yn(p.get("integrated_closed_loop_diagnostic"))),
        ])
    _add_kv_table(doc, active_rows)

    doc.add_paragraph("5. Software specifics (MDSW)", style="MDRH2")
    mdsw_rows: list[tuple[str, str]] = [
        ("Standalone MDSW", _yn(p.get("is_mdsw"))),
    ]
    if p.get("is_mdsw") == "Yes":
        mdsw_rows.extend([
            ("Information drives decisions", _yn(p.get("info_drives_decisions"))),
            ("Decision impact", p.get("decision_significance", "") or "—"),
            ("Monitoring role", p.get("monitoring_role", "") or "—"),
            ("Drives or controls a hardware device", _yn(p.get("drives_hardware_device"))),
        ])
    _add_kv_table(doc, mdsw_rows)

    doc.add_paragraph("6. Special-rule triggers", style="MDRH2")
    trigger_rows: list[tuple[str, str]] = []
    if p.get("rule_14_medicinal_substance"):
        trigger_rows.append(("Rule 14", "Medicinal substance with ancillary action"))
    if p.get("rule_14_blood_derivative"):
        trigger_rows.append(("Rule 14", "Human blood derivative"))
    if p.get("rule_18_non_viable_tissue"):
        trigger_rows.append(("Rule 18", "Non-viable human or animal tissue or derivative"))
    if p.get("rule_15_contraception_or_sti"):
        form = p.get("rule_15_form") or "—"
        trigger_rows.append(("Rule 15", f"Contraception / STI prevention (form: {form})"))
    if p.get("rule_16_disinfection"):
        target = p.get("rule_16_target") or "—"
        trigger_rows.append(("Rule 16", f"Disinfection / cleaning / sterilising (target: {target})"))
    if p.get("rule_17_xray_images"):
        trigger_rows.append(("Rule 17", "Diagnostic images from X-ray"))
    if p.get("rule_19_nanomaterial"):
        exp = p.get("nanomaterial_exposure_potential") or "—"
        trigger_rows.append(("Rule 19", f"Nanomaterial (internal exposure: {exp})"))
    if p.get("rule_20_inhalation"):
        ess = _yn(p.get("rule_20_essential_to_efficacy"))
        trigger_rows.append((
            "Rule 20",
            f"Administers medicines by inhalation (essential to efficacy / life-threatening: {ess})",
        ))
    if p.get("rule_21_absorbed_substance"):
        mode = p.get("rule_21_action_mode") or "—"
        trigger_rows.append(("Rule 21", f"Substances via orifice or skin (mode: {mode})"))
    if trigger_rows:
        _add_kv_table(doc, trigger_rows)
    else:
        doc.add_paragraph("None flagged.", style="MDRBody")

    doc.add_paragraph("7. Use context", style="MDRH2")
    _add_kv_table(doc, [
        ("Intended user", p.get("user_type", "")),
        ("Use environment", p.get("use_environment", "")),
        ("Most critical intended use",
         p.get("multiple_intended_uses") or "—"),
    ])


def _render_classification_report(doc, result: dict, device_params: dict) -> None:
    from docx.shared import Pt, RGBColor

    v4 = result.get("v4_output", {}) or {}

    doc.add_paragraph("MDR Annex VIII Classification Report", style="MDRTitle")
    ts = datetime.now().strftime("%B %d, %Y at %H:%M")
    doc.add_paragraph(ts, style="MDRSubtitle")

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    run = rule.add_run("─" * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = RGBColor(*_TEAL)

    _render_device_parameters(doc, device_params)

    if "error" in v4:
        doc.add_paragraph("Classification Error", style="MDRH1")
        doc.add_paragraph(str(v4.get("error", "")), style="MDRBody")
        return

    s1 = v4.get("section_1_final_mdr_device_classification", {}) or {}
    s2 = v4.get("section_2_primary_mdr_classification_rule", {}) or {}
    s3 = v4.get("section_3_classification_rationale", {}) or {}
    s6 = v4.get("section_6_rule_by_rule_assessment", []) or []
    s7 = v4.get("section_7_applicable_rules", []) or []
    s8 = v4.get("section_8_excluded_rules", {}) or {}
    s9 = v4.get("section_9_rule_conflict_and_precedence_assessment", {}) or {}
    s10 = v4.get("section_10_validation_against_mdr_annex_viii_and_mdcg_2021_24", {}) or {}
    s11 = v4.get("section_11_validation_status", {}) or {}

    doc.add_paragraph(
        f"Provisional Classification: Class {s1.get('provisional_class', '—')}",
        style="MDRH1",
    )
    summary = s1.get("single_sentence_statement", "")
    if summary:
        _add_inline(doc.add_paragraph(style="MDRBody"), summary)
    doc.add_paragraph(
        "This classification is provisional and requires human review under MDR Article 51.",
        style="MDRQuote",
    )

    doc.add_paragraph("Primary Rule", style="MDRH2")
    primary = f"{s2.get('rule', '—')} — {s2.get('sub_clause_description', '—')}"
    _add_inline(doc.add_paragraph(style="MDRBody"), f"**{primary}**")
    statement = s2.get("single_paragraph_statement", "")
    if statement:
        _add_inline(doc.add_paragraph(style="MDRBody"), statement)

    if s3:
        doc.add_paragraph("Classification Rationale", style="MDRH1")
        controlling = s3.get("controlling_logic_prose", "")
        if controlling:
            doc.add_paragraph("Controlling logic", style="MDRH3")
            _add_inline(doc.add_paragraph(style="MDRBody"), controlling)
        verbatim = s3.get("annex_viii_text_verbatim", {}) or {}
        if verbatim.get("text"):
            doc.add_paragraph("Annex VIII text (verbatim)", style="MDRH3")
            _add_inline(doc.add_paragraph(style="MDRQuote"), verbatim["text"])
            src = verbatim.get("source_url", "")
            if src:
                doc.add_paragraph(f"Source: {src}", style="MDRQuote")
        consequence = s3.get("consequence_assessment_prose", "")
        if consequence:
            doc.add_paragraph("Consequence assessment", style="MDRH3")
            _add_inline(doc.add_paragraph(style="MDRBody"), consequence)

    if s6:
        doc.add_paragraph("Rule-by-Rule Assessment", style="MDRH1")
        table = doc.add_table(rows=len(s6) + 1, cols=4)
        table.style = "Light Grid Accent 1"
        header = ["Rule", "Applicability", "Class", "Assessment"]
        for ci, h in enumerate(header):
            table.cell(0, ci).text = h
            for r in table.cell(0, ci).paragraphs[0].runs:
                r.bold = True
        for ri, item in enumerate(s6, 1):
            table.cell(ri, 0).text = item.get("rule", "")
            table.cell(ri, 1).text = item.get("applicability", "")
            table.cell(ri, 2).text = str(item.get("resulting_class_if_applicable") or "")
            table.cell(ri, 3).text = item.get("assessment_prose", "")
        doc.add_paragraph()

    if s7:
        doc.add_paragraph("Applicable Rules", style="MDRH2")
        _add_inline(doc.add_paragraph(style="MDRBody"), ", ".join(f"**{r}**" for r in s7))
    if s8.get("excluded_rules_prose"):
        doc.add_paragraph("Excluded Rules Reasoning", style="MDRH3")
        _add_inline(doc.add_paragraph(style="MDRBody"), s8["excluded_rules_prose"])

    comparable = s10.get("comparable_devices_assessed", []) or []
    if comparable:
        doc.add_paragraph("Comparable Devices Assessed", style="MDRH1")
        for i, dev in enumerate(comparable, 1):
            doc.add_paragraph(
                f"{i}. {dev.get('device_identifier', 'Unknown')}", style="MDRH3"
            )
            score = dev.get("similarity_score_final")
            quality = dev.get("match_quality_rating", "—")
            if score is not None:
                _add_inline(
                    doc.add_paragraph(style="MDRBody"),
                    f"**Similarity:** {score:.2f}   **Quality:** {quality}",
                )
            for label, key in (
                ("Why this comparator scored highly", "precedent_weight_reasoning_prose"),
                ("What transfers", "transfers_reasoning_to_present_prose"),
                ("Limits of transfer", "limits_of_transfer_prose"),
            ):
                txt = dev.get(key, "")
                if txt:
                    _add_inline(
                        doc.add_paragraph(style="MDRBody"), f"**{label}:** {txt}"
                    )
            purpose = dev.get("intended_purpose_verbatim_quote", "")
            if purpose:
                _add_inline(doc.add_paragraph(style="MDRQuote"), purpose)
                src = dev.get("intended_purpose_quote_url", "")
                if src:
                    doc.add_paragraph(f"Source: {src}", style="MDRQuote")

    uncertainties = s9.get("interpretive_uncertainties_flagged", []) or []
    if uncertainties:
        doc.add_paragraph("Interpretive Uncertainties", style="MDRH1")
        if s9.get("precedence_prose"):
            _add_inline(doc.add_paragraph(style="MDRBody"), s9["precedence_prose"])
        for unc in uncertainties:
            tag = "(load-bearing) " if unc.get("load_bearing_to_classification") else ""
            doc.add_paragraph(
                f"{tag}{unc.get('uncertainty_type', '')}", style="MDRH3"
            )
            if unc.get("substrate_provisional_position_prose"):
                _add_inline(
                    doc.add_paragraph(style="MDRBody"),
                    f"*Provisional position:* {unc['substrate_provisional_position_prose']}",
                )
            if unc.get("human_deliberation_rationale_prose"):
                _add_inline(
                    doc.add_paragraph(style="MDRBody"),
                    f"*Human deliberation needed:* {unc['human_deliberation_rationale_prose']}",
                )

    checks = s10.get("validation_checks", {}) or {}
    if checks or s10.get("validation_summary_prose") or s11:
        doc.add_paragraph("Validation", style="MDRH1")
        for name, passed in checks.items():
            marker = "OK" if passed else "FAIL"
            _add_inline(
                doc.add_paragraph(style="MDRBody"),
                f"**[{marker}]** {name.replace('_', ' ')}",
            )
        if s10.get("validation_summary_prose"):
            _add_inline(doc.add_paragraph(style="MDRBody"), s10["validation_summary_prose"])
        if s11.get("status"):
            _add_inline(
                doc.add_paragraph(style="MDRBody"),
                f"**Status:** {s11['status']}",
            )
        if s11.get("closing_statement_prose"):
            _add_inline(doc.add_paragraph(style="MDRBody"), s11["closing_statement_prose"])

    doc.add_paragraph(
        "This report was generated by the MDR Classification pipeline. "
        "All outputs are provisional and require qualified human review.",
        style="MDRDisclaimer",
    )


def generate_docx_bytes(result: dict, device_params: dict) -> bytes:
    """Return a formatted .docx report of the classification."""
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    _add_styles(doc)
    _render_classification_report(doc, result, device_params)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_bundle_zip(result: dict, device_params: dict) -> tuple[bytes, str]:
    """Bundle the Word document, individual JSON packages and the audit log
    into a single ZIP.

    Returns (zip_bytes, filename).
    """
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    audit_log = result.get("audit_log") or {}

    docx_bytes = generate_docx_bytes(result, device_params)

    full_session = {
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "device_parameters": device_params,
        "v4_output": result.get("v4_output"),
        "regulatory_raw": result.get("regulatory_raw"),
        "comparator_raw": result.get("comparator_raw"),
        "orchestrator_raw": result.get("orchestrator_raw"),
        "audit_log": audit_log,
    }

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"classification_{timestamp}.docx", docx_bytes)
        zf.writestr(
            "full_session.json",
            json.dumps(full_session, indent=2, ensure_ascii=False, default=str),
        )
        zf.writestr(
            "v4_classification.json",
            json.dumps(result.get("v4_output") or {}, indent=2, ensure_ascii=False, default=str),
        )
        zf.writestr(
            "regulatory_search.json",
            json.dumps(result.get("regulatory_raw") or {}, indent=2, ensure_ascii=False, default=str),
        )
        zf.writestr(
            "comparator_engine.json",
            json.dumps(result.get("comparator_raw") or {}, indent=2, ensure_ascii=False, default=str),
        )
        zf.writestr(
            "audit_log.json",
            json.dumps(audit_log, indent=2, ensure_ascii=False, default=str),
        )

    filename = f"mdr_classification_{timestamp}.zip"
    return buf.getvalue(), filename
